"""Erzeugt die Präsentation „EMRIQ Kelmis – Leistungskonzept“ auf Basis der Drees-&-Sommer-Vorlage.

Aufruf: python3 praesentation/build_pptx.py
Abbildungen liegen als SVG in praesentation/svg und werden mit PNG-Ersatzbild eingebettet
(PowerPoint zeigt die Vektorgrafik, ältere Programme das PNG).
"""
import copy
import os
import subprocess

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
TEMPLATE = os.path.join(HERE, "vorlage", "Vorlage.pptx")
SVG = os.environ.get("EMRIQ_SVG", os.path.join(HERE, "svg"))
PNG = os.environ.get("EMRIQ_PNG", os.path.join(HERE, "png"))
OUT = os.environ.get("EMRIQ_OUT", os.path.join(ROOT, "export", "EMRIQ_Kelmis_Leistungskonzept.pptx"))

FOOTER = "EMRIQ Kelmis – Drees & Sommer – 14.10.2026"
NAVY, BLUE, LBLUE, GREEN = "002965", "005AA7", "009FE3", "95C11F"
TXT, GREY1, GREY2, LINE, WHITE = "404040", "F2F2F2", "E6E6E6", "CDCDCD", "FFFFFF"

os.makedirs(PNG, exist_ok=True)
prs = Presentation(TEMPLATE)
LAYOUT = {l.name.strip(): l for l in prs.slide_layouts}
svg_counter = [0]
svg_parts = {}


# ---------------------------------------------------------------- Hilfsfunktionen
def rgb(h):
    return RGBColor.from_string(h)


def render_png(svg_path, color_swap=None):
    """Rendert ein SVG (optional umgefärbt) als PNG und gibt (svg_bytes, png_path) zurück."""
    data = open(svg_path, encoding="utf-8").read()
    name = os.path.splitext(os.path.basename(svg_path))[0]
    if color_swap:
        for a, b in color_swap.items():
            data = data.replace(a, b)
        name += "_" + "_".join(v.strip("#") for v in color_swap.values())
    tmp_svg = os.path.join(PNG, name + ".svg")
    png = os.path.join(PNG, name + ".png")
    with open(tmp_svg, "w", encoding="utf-8") as f:
        f.write(data)
    if not os.path.exists(png) or os.path.getmtime(png) < os.path.getmtime(tmp_svg):
        subprocess.run(["node", os.path.join(HERE, "render_svg.js"), tmp_svg, png, "3"], check=True, capture_output=True)
    return data.encode("utf-8"), png


def add_svg(slide, svg_path, x, y, w, h, name, color_swap=None):
    blob, png = render_png(svg_path, color_swap)
    pic = slide.shapes.add_picture(png, Inches(x), Inches(y), Inches(w), Inches(h))
    pic.name = name
    key = (svg_path, tuple(sorted((color_swap or {}).items())))
    if key not in svg_parts:
        svg_counter[0] += 1
        svg_parts[key] = Part(PackURI(f"/ppt/media/emriq_svg{svg_counter[0]}.svg"), "image/svg+xml",
                              slide.part.package, blob)
    rid = slide.part.relate_to(svg_parts[key], RT.IMAGE)
    blip = pic._element.find(".//" + qn("a:blip"))
    ext_lst = etree.SubElement(blip, qn("a:extLst"))
    ext = etree.SubElement(ext_lst, qn("a:ext"), uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
    svg_blip = etree.SubElement(ext, "{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip",
                                nsmap={"asvg": "http://schemas.microsoft.com/office/drawing/2016/SVG/main"})
    svg_blip.set(qn("r:embed"), rid)
    return pic


def icon(slide, name, x, y, d, circle=BLUE, stroke="#FFFFFF"):
    """Linien-Icon im Farbkreis (Durchmesser d in Zoll)."""
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.name = f"Icon-Kreis {name}"
    fill(c, circle)
    pad = d * 0.2
    add_svg(slide, os.path.join(SVG, "icons", name + ".svg"), x + pad, y + pad, d - 2 * pad, d - 2 * pad,
            f"Icon {name}", color_swap={"#002965": stroke})


def fill(shape, color=None, line=None):
    if color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(color)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False


def box(slide, x, y, w, h, color=GREY1, name="Fläche", line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.name = name
    fill(s, color, line)
    s.text_frame.text = ""
    return s


def _bullet(p, color):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(0.2))))
    pPr.set("indent", str(-Emu(Inches(0.2))))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    buClr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(buClr, qn("a:srgbClr"), val=color)
    etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
    etree.SubElement(pPr, qn("a:buChar"), char="•")


def text(slide, x, y, w, h, paras, size=14, color=TXT, anchor="t", name="Text", margin=0.0, align="l"):
    """paras: Liste aus Strings oder Dicts {t, size, bold, color, bullet, before, after, runs}."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    m = Inches(margin)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    for i, para in enumerate(paras):
        if isinstance(para, str):
            para = {"t": para}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[para.get("align", align)]
        p.space_before = Pt(para.get("before", 0))
        p.space_after = Pt(para.get("after", 4))
        runs = para.get("runs") or [{"t": para["t"]}]
        for r in runs:
            run = p.add_run()
            run.text = r["t"]
            f = run.font
            f.size = Pt(r.get("size", para.get("size", size)))
            f.bold = r.get("bold", para.get("bold", False))
            f.color.rgb = rgb(r.get("color", para.get("color", color)))
        if para.get("bullet"):
            _bullet(p, para.get("bullet_color", BLUE))
    return tb


def next_id(slide):
    ids = [int(e.get("id")) for e in slide._element.iter() if e.tag == qn("p:cNvPr") and e.get("id", "").isdigit()]
    return max(ids + [1]) + 1


def footer(slide):
    """Fußzeile und Foliennummer wie in der Vorlage (Platzhalter des Layouts)."""
    tree = slide.shapes._spTree
    for kind, idx, name in (("sldNum", 12, "Foliennummernplatzhalter"), ("ftr", 11, "Fußzeilenplatzhalter")):
        if kind == "sldNum":
            body = ('<a:p><a:fld id="{9829449A-3FF3-45CE-9F4F-09BADC0910C6}" type="slidenum"><a:rPr lang="de-DE"/>'
                    '<a:t>‹Nr.›</a:t></a:fld><a:endParaRPr lang="de-DE"/></a:p>')
        else:
            body = f'<a:p><a:r><a:rPr lang="de-DE"/><a:t>{FOOTER.replace("&", "&amp;")}</a:t></a:r></a:p>'
        xml = (f'<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
               f'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
               f'<p:nvSpPr><p:cNvPr id="{next_id(slide)}" name="{name}"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
               f'<p:nvPr><p:ph type="{kind}" sz="quarter" idx="{idx}"/></p:nvPr></p:nvSpPr><p:spPr/>'
               f'<p:txBody><a:bodyPr/><a:lstStyle/>{body}</p:txBody></p:sp>')
        tree.append(etree.fromstring(xml))


def new_slide(title, subtitle, notes):
    s = prs.slides.add_slide(LAYOUT["Only Title"])
    for ph in s.placeholders:
        if ph.placeholder_format.type is not None and ph.placeholder_format.idx == 0:
            ph.text_frame.text = title
        elif ph.placeholder_format.idx == 14:
            ph.text_frame.text = subtitle
    footer(s)
    s.notes_slide.notes_text_frame.text = notes
    return s


def ap_tag(slide, x, y, label, color=NAVY, w=0.62):
    t = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.27))
    t.name = f"Etikett {label}"
    fill(t, color)
    tf = t.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = rgb(WHITE)
    return t


# ---------------------------------------------------------------- Vorlage vorbereiten
# Musterfolien 2 und 3 entfernen, Titelfolie behalten
sld_ids = prs.slides._sldIdLst
for sld_id in list(sld_ids)[1:]:
    prs.part.drop_rel(sld_id.rId)
    sld_ids.remove(sld_id)

# ---------------------------------------------------------------- Folie 1: Titel
s1 = prs.slides[0]
for sp in list(s1.placeholders):
    pf = sp.placeholder_format
    if pf.idx == 15:  # leerer Bildplatzhalter
        sp._element.getparent().remove(sp._element)
    elif pf.idx == 0:
        sp.text_frame.paragraphs[0].runs[0].text = "EMRIQ Kelmis"
    elif pf.idx == 1:
        ps = sp.text_frame.paragraphs
        ps[0].runs[0].text = "Städtebauliches Entwicklungskonzept mit Umsetzungsplan – Leistungskonzept"
        ps[1].runs[0].text = "Ministerium der Deutschsprachigen Gemeinschaft  ·  Fachbereich Raumordnung, Wohnungswesen und Energie"
        ps[2].runs[0].text = "Gesprächstermin  ·  14. Oktober 2026"
    elif pf.idx == 11:
        sp.text_frame.paragraphs[0].runs[0].text = FOOTER
s1.notes_slide.notes_text_frame.text = (
    "Begrüßung und Vorstellung. Ziel des Termins: unser Verständnis der Aufgabe und das Leistungskonzept vorstellen, "
    "Fragen des Ministeriums aufnehmen und offene Punkte für das Angebot klären. Dauer rund 30 Minuten Vortrag, "
    "danach Diskussion.")

# ---------------------------------------------------------------- Folie 2: Ausgangslage
s = new_slide("Ausgangslage und Aufgabe",
              "Ein Innovationsquartier mit klarem Profil – verankert in Kelmis und der Euregio",
              "Kelmis liegt zwischen den Hochschulstandorten Aachen, Lüttich und Maastricht, rund 8 km von Aachen, "
              "mit Bahnhalt Hergenrath und Buslinie 24. Das Kandidatengebiet für das Einstein-Teleskop grenzt an. "
              "Die DG will hier ein thematisch profiliertes Quartier, kein Gewerbegebiet. Die sechs Erfolgskriterien "
              "rechts sind unser Maßstab für die gesamte Bearbeitung.")
add_svg(s, os.path.join(SVG, "01_euregio_lage.svg"), 0.57, 1.75, 5.896, 4.896, "Lagekarte Euregio")
text(s, 6.95, 1.80, 5.8, 1.55, [
    {"t": "Was die Deutschsprachige Gemeinschaft beabsichtigt", "size": 16, "bold": True, "color": NAVY, "after": 6},
    {"t": "Ein Quartier, das Forschung, innovative Unternehmen, Produktion, Gründungen, Wohnen und Lebensqualität "
          "verbindet und die Brückenfunktion der DG zwischen Belgien, Nordrhein-Westfalen und den Hochschulen in "
          "Aachen, Lüttich und Maastricht stärkt."},
], name="Text Absicht")
text(s, 6.95, 3.50, 5.8, 0.35, [{"t": "Woran wir den Erfolg messen", "size": 16, "bold": True, "color": NAVY}],
     name="Überschrift Erfolgskriterien")
crit = ["Schlüssige thematische Profilierung mit Alleinstellung in der Euregio",
        "Nachgewiesene Bedarfe und klar beschriebene Zielgruppen",
        "Robustes Struktur- und Nutzungskonzept mit Entwicklungsspielräumen",
        "Wirtschaftlich und organisatorisch tragfähiges Umsetzungsmodell",
        "Priorisierte Roadmap mit Kosten, Folgekosten und Zuständigkeiten",
        "Standortstory in Deutsch, Französisch und Englisch"]
for i, c in enumerate(crit):
    y = 3.95 + i * 0.44
    add_svg(s, os.path.join(SVG, "icons", "haken.svg"), 6.95, y + 0.02, 0.26, 0.26, f"Haken {i + 1}",
            color_swap={"#002965": "#" + GREEN})
    text(s, 7.35, y, 5.4, 0.36, [c], name=f"Kriterium {i + 1}", anchor="m")

# ---------------------------------------------------------------- Folie 3: Ansatz
s = new_slide("Unser Ansatz",
              "Strategie, Planung und Umsetzung in einem Prozess – aufbauend auf den vorhandenen Grundlagen",
              "Vier Grundsätze prägen unser Vorgehen. Wir nutzen die vorhandenen Studien als Startpunkt und investieren "
              "das Budget in Fokussierung, Zahlen und Umsetzung. Die thematische Entscheidung bleibt bei der DG; wir "
              "liefern eine transparente Entscheidungsgrundlage. Besonders ist die wirtschaftliche Betrachtung aus "
              "Sicht von Eigentümer, Käufern und Investoren.")
cards = [
    ("grundlagen", "Grundlagen nutzen",
     "Fachgutachten Raumordnung, Gewerbeflächenstrategie, OBL'EU und MORO werden zu Beginn zusammengeführt. "
     "Wir ergänzen gezielt, statt neu zu analysieren."),
    ("fokus", "Fokussieren",
     "Ein Thema statt Campusstrategie. Schwerpunkte werden transparent bewertet; die DG entscheidet auf "
     "nachvollziehbarer Grundlage."),
    ("rechnen", "Rechnen",
     "Das Quartier wird aus Sicht von Eigentümer, Käufern und Investoren durchgerechnet: Exit-Szenarien, "
     "Grundstückswerte, Finanzierungsbedarf."),
    ("umsetzen", "Umsetzen und vermitteln",
     "Roadmap mit Zuständigkeiten und Meilensteinen sowie eine dreisprachige Standortstory für Investoren "
     "und Fachkräfte."),
]
cw, gap = 2.82, 0.30
for i, (ic, head, body) in enumerate(cards):
    x = 0.57 + i * (cw + gap)
    box(s, x, 1.85, cw, 3.45, GREY1, name=f"Karte {head}")
    icon(s, ic, x + 0.25, 2.10, 0.75)
    text(s, x + 0.25, 2.95, cw - 0.5, 0.55, [{"t": head, "size": 16, "bold": True, "color": NAVY}], anchor="b",
         name=f"Titel {head}")
    text(s, x + 0.25, 3.62, cw - 0.5, 1.6, [body], name=f"Text {head}")
band = box(s, 0.57, 5.55, 12.18, 1.05, NAVY, name="Team")
icon(s, "team", 0.82, 5.70, 0.75, circle=BLUE)
text(s, 1.85, 5.62, 10.7, 0.92, [
    {"t": "Interdisziplinäres Kernteam: Stadt- und Quartiersentwicklung · Immobilienökonomie · Projektmanagement · "
          "Nachhaltigkeit · Mobilität · Kommunikation", "bold": True, "color": WHITE, "after": 4},
    {"t": "Projektleitung und Team mit Deutsch, Französisch und Englisch · persönlicher Ansprechpartner, "
          "bei Bedarf vor Ort in Eupen", "color": WHITE},
], anchor="m", name="Text Team")

# ---------------------------------------------------------------- Folie 4: Leistungsumfang
s = new_slide("Leistungsumfang im Überblick",
              "Acht Arbeitspakete in drei Phasen, begleitet von einem durchgängigen Abstimmungsprozess",
              "Die Arbeitspakete folgen der Logik des Lastenhefts: erst Grundlagen und Analyse, dann Profil und "
              "Konzept, zuletzt Wirtschaftlichkeit und Umsetzung. AP 8 begleitet alle Phasen. Jedes Paket endet mit "
              "einem prüfbaren Ergebnis und einem Meilenstein mit Freigabe durch die DG.")
add_svg(s, os.path.join(SVG, "04_leistungsarchitektur.svg"), 0.57, 1.75, 12.177, 4.854, "Leistungsarchitektur")

# ---------------------------------------------------------------- Folie 5: AP 1–3
s = new_slide("AP 1–3: Grundlagen, Analyse und Themenprofil",
              "Vom gemeinsamen Zielsystem zur begründeten thematischen Ausrichtung",
              "AP 1 schafft eine gemeinsame Arbeitsgrundlage und einen Kriterienrahmen. AP 2 analysiert Standort, "
              "Mobilität und Markt einschließlich möglicher Effekte des Einstein-Teleskops. In AP 3 wird das Thema "
              "in vier Stufen eingegrenzt; die Entscheidung fällt im ersten Workshop. Rückfrage an das Ministerium: "
              "Gibt es bereits eine politische Themenpräferenz?")
aps = [
    ("AP 1", "Projektstart, Grundlagen und Zielsystem",
     ["Auswertung vorhandener Gutachten, Strategien und Geodaten",
      "Zielsystem, Betrachtungsraum und Kriterienrahmen"],
     "Abgestimmte Arbeitsgrundlage und Bearbeitungsplan"),
    ("AP 2", "Standort-, Markt- und Potenzialanalyse",
     ["Standort, Mobilität und Wohnungsmarkt",
      "Vergleichsstandorte, Nachfrage, Effekte des Einstein-Teleskops"],
     "Standortsteckbrief mit SWOT und Entwicklungshebeln"),
    ("AP 3", "Themenprofil, Zielgruppen und Ecosystem",
     ["Bewertung der Schwerpunkte, Empfehlung mit Alternativen",
      "Zielgruppen, Anforderungen, Partner- und Betreibermodell"],
     "Begründetes Themen- und Zielgruppenprofil"),
]
for i, (tag, head, bullets, result) in enumerate(aps):
    y = 1.80 + i * 1.62
    ap_tag(s, 0.57, y + 0.02, tag)
    text(s, 1.35, y, 5.4, 0.32, [{"t": head, "size": 15, "bold": True, "color": NAVY}], name=f"Titel {tag}")
    text(s, 1.35, y + 0.36, 5.4, 0.78, [{"t": b, "bullet": True, "after": 2} for b in bullets], name=f"Inhalt {tag}")
    text(s, 1.35, y + 1.12, 5.4, 0.3, [{"runs": [{"t": "Ergebnis: ", "bold": True, "color": GREEN},
                                                 {"t": result}]}], name=f"Ergebnis {tag}")
add_svg(s, os.path.join(SVG, "05_themenprofil_trichter.svg"), 7.05, 1.75, 5.698, 4.854, "Trichter Themenprofil")

# ---------------------------------------------------------------- Folie 6: AP 4–5
s = new_slide("AP 4–5: Entwicklungskonzept, Nachhaltigkeit und Mobilität",
              "Varianten, Vorzugsvariante und Phasierung – verknüpft mit Ortskern und Bahnhof Hergenrath",
              "Das Raumbild zeigt die Suchräume aus den Vorarbeiten; der Betrachtungsraum wird in AP 1 gemeinsam "
              "festgelegt. Entscheidend ist die Verbindung zum Bahnhof Hergenrath, weil die Erreichbarkeit mit Bus und "
              "Bahn heute der größte Engpass ist. Die Vorzugsvariante wird mit der DG verbindlich festgelegt und ist "
              "Grundlage für Wirtschaftlichkeit und Umsetzungsplan.")
add_svg(s, os.path.join(SVG, "06_raumschema.svg"), 0.57, 1.75, 6.396, 4.896, "Raumschema Kelmis")
panels = [
    ("AP 4", "Städtebauliches Entwicklungs- und Nutzungskonzept",
     ["Räumlich-funktionales Leitbild und Flächenprogramm", "Städtebau, Freiraum, Erschließung, Infrastruktur",
      "Variantenvergleich und Vorzugsvariante", "Phasierung, Flexibilität, erste Aktivierungsmaßnahmen"]),
    ("AP 5", "Nachhaltigkeit, Mobilität und Resilienz",
     ["Ziele für Energie, Klima, Ressourcen, Biodiversität und soziale Qualität",
      "Multimodale Erreichbarkeit mit weniger Autoverkehr", "Klimaangepasste Freiräume, Regenwasser, zirkuläres Bauen",
      "Messbare Indikatoren für Planung und Umsetzung"]),
]
for i, (tag, head, bullets) in enumerate(panels):
    y = 1.80 + i * 2.45
    box(s, 7.25, y, 5.5, 2.30, GREY1, name=f"Fläche {tag}")
    ap_tag(s, 7.45, y + 0.2, tag)
    text(s, 8.20, y + 0.16, 4.4, 0.62, [{"t": head, "size": 15, "bold": True, "color": NAVY}], name=f"Titel {tag}")
    text(s, 7.45, y + 0.80, 5.15, 1.45, [{"t": b, "bullet": True, "after": 3} for b in bullets], name=f"Inhalt {tag}")

# ---------------------------------------------------------------- Folie 7: AP 6
s = new_slide("AP 6: Wirtschaftlichkeit, Business Case und Finanzierung",
              "Das Quartier aus Sicht von Eigentümer, Käufern und Investoren rechnen",
              "Basisstrategie ist die Entwicklung bis zum Planungsrecht mit anschließender Erschließung und parallelem "
              "Verkauf der Grundstücke. So können Erlöse früh fließen und einen Teil der Vorleistungen refinanzieren. "
              "Die Grundstückspreise schätzen wir über die Residualwertmethode aus Sicht der späteren Projektentwickler. "
              "Wir vergleichen verschiedene Exit-Zeitpunkte und empfehlen Grundstücksgrößen, die für Käufer "
              "finanzierbar sind. Rechtliche Bewertungen und Verkehrswertgutachten sind nicht Teil der Leistung.")
text(s, 0.57, 1.78, 6.0, 0.3, [{"t": "Basisstrategie und Exit-Zeitpunkte", "size": 15, "bold": True, "color": NAVY}],
     name="Überschrift Exit")
add_svg(s, os.path.join(SVG, "07_wertentwicklung_exit.svg"), 0.57, 2.12, 6.0, 3.198, "Wertentwicklung Exit")
text(s, 6.85, 1.78, 5.9, 0.3, [{"t": "Grundstückswert nach Residualwertmethode", "size": 15, "bold": True,
                                "color": NAVY}], name="Überschrift Residualwert")
add_svg(s, os.path.join(SVG, "07_residualwert.svg"), 6.85, 2.12, 5.896, 3.198, "Residualwert")
cards = [("Exit-Szenarien", "Verkauf nach Planungsrecht, parallel zur Erschließung oder nach Vermarktungsstart"),
         ("Slot-Sizes", "Grundstücks- und Investitionsgrößen, die finanzierbar und kapitalmarktfähig sind"),
         ("Parzellierung", "Eigenständig erschließbare, entwickelbare und vermarktbare Grundstücke"),
         ("Öffentliche Hand", "Kosten, Erlöse und Finanzierungsbedarf, Folgekosten und Einnahmen")]
for i, (head, body) in enumerate(cards):
    x = 0.57 + i * (cw + gap)
    box(s, x, 5.42, cw, 1.22, GREY1, name=f"Karte {head}")
    text(s, x + 0.18, 5.50, cw - 0.36, 1.08, [{"t": head, "size": 14, "bold": True, "color": NAVY, "after": 2},
                                              {"t": body, "size": 13}], name=f"Text {head}")

# ---------------------------------------------------------------- Folie 8: AP 7–8
s = new_slide("AP 7–8: Umsetzung, Kommunikation und Vermarktung",
              "Priorisierte Roadmap und schlanke, verbindliche Abstimmungsformate",
              "AP 7 übersetzt die Ergebnisse in einen Umsetzungsplan mit zehn priorisierten Maßnahmen, Zuständigkeiten "
              "und Risiken. AP 8 begleitet die gesamte Bearbeitung mit festen Formaten. Akteure vor Ort und in der "
              "Euregio binden wir nur nach Freigabe durch die DG ein. Die Vermarktungsunterlagen entstehen nach AP 6 "
              "und AP 7 und werden fachlich, nicht nur sprachlich, ins Französische und Englische übertragen.")
text(s, 0.57, 1.78, 5.9, 0.3, [{"runs": [{"t": "AP 7  ", "bold": True, "color": BLUE},
                                         {"t": "Umsetzungsstrategie und Roadmap", "bold": True, "color": NAVY}],
                                "size": 15}], name="Überschrift AP 7")
rows = [("massnahmen", "Top-10-Maßnahmenkatalog mit Prioritäten, Abhängigkeiten und Entscheidungsbedarfen"),
        ("termine", "Termin- und Meilensteinplanung mit Abschluss deutlich vor dem 31.10.2028"),
        ("verantwortung", "Zuständigkeiten, Ressourcen und Schnittstellen"),
        ("verfahren", "Genehmigungs-, Planungs- und Vergabefahrplan auf Verfahrensebene"),
        ("risiko", "Risikoanalyse mit Gegenmaßnahmen")]
for i, (ic, t) in enumerate(rows):
    y = 2.25 + i * 0.86
    icon(s, ic, 0.57, y, 0.62)
    text(s, 1.40, y, 4.95, 0.62, [t], anchor="m", name=f"AP 7 Punkt {i + 1}")
text(s, 6.85, 1.78, 5.9, 0.3, [{"runs": [{"t": "AP 8  ", "bold": True, "color": BLUE},
                                         {"t": "Kommunikation, Beteiligung und Vermarktung", "bold": True,
                                          "color": NAVY}], "size": 15}], name="Überschrift AP 8")
stats = [("1", "Auftakt", "digital, bis 90 Minuten"), ("4", "Abstimmungen", "digital, je bis 45 Minuten"),
         ("2", "Workshops", "je bis 3 Stunden, einer in Präsenz"), ("1", "Abschluss", "Präsentation, digital, bis 90 Minuten")]
for i, (num, head, sub) in enumerate(stats):
    x = 6.85 + (i % 2) * 3.0
    y = 2.25 + (i // 2) * 1.32
    box(s, x, y, 2.8, 1.15, GREY1, name=f"Format {head}")
    text(s, x + 0.15, y + 0.08, 0.75, 1.0, [{"t": num, "size": 40, "bold": True, "color": NAVY}], anchor="m",
         name=f"Zahl {head}", align="c")
    text(s, x + 0.95, y + 0.1, 1.8, 0.95, [{"t": head, "size": 14, "bold": True, "color": NAVY, "after": 2},
                                           {"t": sub, "size": 13}], anchor="m", name=f"Text {head}")
box(s, 6.85, 4.98, 5.8, 1.62, NAVY, name="Vermarktung")
icon(s, "sprache", 7.07, 5.18, 0.6)
icon(s, "karte", 7.07, 5.88, 0.6)
text(s, 7.85, 5.18, 4.65, 0.6, [{"t": "Standortstory und Vermarktungsbausteine in Deutsch, Französisch und Englisch",
                                 "color": WHITE}], anchor="m", name="Text Sprachen")
text(s, 7.85, 5.88, 4.65, 0.6, [{"t": "GIS-fähige Karten im belgischen Lambert-System, kurzfristige "
                                      "Aktivierungsmaßnahmen", "color": WHITE}], anchor="m", name="Text GIS")

# ---------------------------------------------------------------- Folie 9: Zeitplan
s = new_slide("Zeitplan und Meilensteine",
              "Abschluss nach rund 16 Monaten – deutlich vor dem Endtermin am 31. Oktober 2028",
              "Sieben Meilensteine mit Freigabe durch die DG. Der Terminplan wird im Kick-off konkretisiert. Bei "
              "Beauftragung um den Jahreswechsel liegt der Abschluss im Frühjahr 2028, also mit deutlichem Puffer vor "
              "dem Endtermin. Die beiden Workshops liegen an den Stellen, an denen Entscheidungen der DG gebraucht "
              "werden: Themenprofil und Entwicklungskonzept sowie Wirtschaftlichkeit und Umsetzung.")
add_svg(s, os.path.join(SVG, "09_terminplan.svg"), 0.57, 1.80, 12.177, 4.75, "Terminplan")

# ---------------------------------------------------------------- Folie 10: Liefergegenstände
s = new_slide("Liefergegenstände und nächste Schritte",
              "Was die Deutschsprachige Gemeinschaft erhält und wie wir starten",
              "Alle Unterlagen und Rechte gehen an die DG über. Nicht enthalten sind Planungen über die Konzeptebene "
              "hinaus, Sondergutachten (z. B. Verkehr, Schall, Altlasten, Wertermittlung) und Rechtsberatung; ein "
              "solcher Bedarf wird im Umsetzungsplan mit Zeitpunkt und Kostenansatz ausgewiesen. Zum Abschluss die "
              "Fragen an das Ministerium: Betrachtungsraum, Themenpräferenz, gewünschte Formate der Ergebnisse.")
text(s, 0.57, 1.78, 7.6, 0.3, [{"t": "Liefergegenstände", "size": 15, "bold": True, "color": NAVY}],
     name="Überschrift Liefergegenstände")
items = ["Projekt- und Bearbeitungsplan mit Termin-, Risiko- und Entscheidungslogik",
         "Standort-, Markt- und Potenzialanalyse mit Standortsteckbrief",
         "Themen-, Zielgruppen- und Ecosystem-Konzept",
         "Entwicklungs- und Nutzungskonzept mit Vorzugsvariante",
         "Nachhaltigkeits-, Mobilitäts- und Infrastrukturkonzept",
         "Wirtschaftlichkeitsbetrachtung mit offengelegten Annahmen",
         "Finanzierungs-, Förder-, Träger- und Betreibermodell",
         "Umsetzungsplan mit Maßnahmen, Prioritäten und Zuständigkeiten",
         "Vermarktungsunterlagen in Deutsch, Französisch und Englisch",
         "GIS-Karten, Protokolle, Abschlussbericht mit bearbeitbaren Dateien"]
for i, t in enumerate(items):
    col, row = i // 5, i % 5
    x, y = 0.57 + col * 3.9, 2.22 + row * 0.62
    add_svg(s, os.path.join(SVG, "icons", "haken.svg"), x, y + 0.05, 0.26, 0.26, f"Haken L{i + 1}",
            color_swap={"#002965": "#" + GREEN})
    text(s, x + 0.38, y, 3.35, 0.56, [{"t": t, "size": 13}], name=f"Liefergegenstand {i + 1}")
box(s, 0.57, 5.40, 7.55, 1.2, GREY1, name="Qualitätssicherung")
icon(s, "qualitaet", 0.77, 5.62, 0.75)
text(s, 1.72, 5.48, 6.25, 1.05, [{"t": "Qualitätssicherung", "size": 14, "bold": True, "color": NAVY, "after": 2},
                                 {"t": "Vier-Augen-Prinzip, fachübergreifende Reviews an den Meilensteinen, "
                                       "dokumentierte Freigaben, fachliche Prüfung der französischen und englischen "
                                       "Fassung", "size": 13}], anchor="m", name="Text Qualitätssicherung")
box(s, 8.45, 1.80, 4.3, 4.80, NAVY, name="Nächste Schritte")
text(s, 8.70, 2.00, 3.85, 0.4, [{"t": "Nächste Schritte", "size": 16, "bold": True, "color": WHITE}],
     name="Überschrift Nächste Schritte")
steps = ["Offene Punkte aus dem heutigen Termin klären: Betrachtungsraum, Themenpräferenz, Formate der Ergebnisse",
         "Angebot auf dieser Grundlage abschließen",
         "Beauftragung und Auftakt im ersten Monat",
         "Bereitstellung der Grundlagen: Fachgutachten, Gewerbeflächenstrategie, MORO-Bericht, Geodaten"]
yy = 2.55
for i, t in enumerate(steps):
    n = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.70), Inches(yy), Inches(0.42), Inches(0.42))
    n.name = f"Nummer {i + 1}"
    fill(n, BLUE)
    tf = n.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(i + 1)
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = rgb(WHITE)
    h = 0.95 if len(t) > 60 else 0.62
    text(s, 9.30, yy - 0.02, 3.25, h, [{"t": t, "color": WHITE}], name=f"Schritt {i + 1}")
    yy += h + 0.12

prs.core_properties.title = "EMRIQ Kelmis – Leistungskonzept"
prs.core_properties.author = ""
prs.core_properties.last_modified_by = ""
os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print("geschrieben:", OUT, len(prs.slides._sldIdLst), "Folien")
